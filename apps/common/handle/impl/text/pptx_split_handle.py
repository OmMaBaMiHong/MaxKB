# coding=utf-8
"""
    PPTX 解析 handler（M2）：python-pptx 逐页提取形状文本与演讲者备注，
    页与页之间以二级标题分隔，交给上游 SplitModel 切段。
"""
from typing import List

from common.handle.base_split_handle import BaseSplitHandle
from common.handle.impl.text.pdf_split_handle import default_pattern_list
from common.utils.split_model import SplitModel


class PptxSplitHandle(BaseSplitHandle):
    def support(self, file, get_buffer):
        file_name: str = file.name.lower()
        return file_name.endswith(".pptx")

    def handle(self, file, pattern_list: List, with_filter: bool, limit: int, get_buffer, save_image):
        content = self.get_content(file, save_image)
        if pattern_list is not None and len(pattern_list) > 0:
            split_model = SplitModel(pattern_list, with_filter, limit)
        else:
            split_model = SplitModel(default_pattern_list, with_filter=with_filter, limit=limit)
        return {"name": file.name, "content": split_model.parse(content)}

    def get_content(self, file, save_image):
        import io
        import tempfile

        from pptx import Presentation

        with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as temp_file:
            for chunk in file.chunks():
                temp_file.write(chunk)
            temp_path = temp_file.name
        try:
            presentation = Presentation(temp_path)
            parts = []
            for index, slide in enumerate(presentation.slides, start=1):
                texts = []
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        texts.append(shape.text_frame.text.strip())
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                    texts.append(slide.notes_slide.notes_text_frame.text.strip())
                if texts:
                    parts.append(f"## 第 {index} 页\n\n" + "\n\n".join(texts))
            return "\n\n".join(parts)
        finally:
            import os

            os.remove(temp_path)
